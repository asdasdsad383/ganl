#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
分析论文格式模板文件
"""
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

def analyze_docx_format(file_path):
    """分析docx文件的格式"""
    doc = Document(file_path)

    print("=" * 60)
    print("论文格式模板分析报告")
    print("=" * 60)

    # 分析页面设置
    section = doc.sections[0]
    print("\n【页面设置】")
    print(f"页边距:")
    print(f"  上: {section.top_margin.cm:.2f} cm")
    print(f"  下: {section.bottom_margin.cm:.2f} cm")
    print(f"  左: {section.left_margin.cm:.2f} cm")
    print(f"  右: {section.right_margin.cm:.2f} cm")
    print(f"页面尺寸: {section.page_width.cm:.2f} x {section.page_height.cm:.2f} cm")

    # 分析样式
    print("\n【样式分析】")
    styles_used = {}

    for para in doc.paragraphs:
        if para.text.strip():  # 只分析非空段落
            style_name = para.style.name

            if style_name not in styles_used:
                styles_used[style_name] = {
                    'examples': [],
                    'font_name': None,
                    'font_size': None,
                    'alignment': None,
                    'line_spacing': None,
                    'space_before': None,
                    'space_after': None,
                    'first_line_indent': None,
                    'bold': None,
                }

            # 保存示例文本
            if len(styles_used[style_name]['examples']) < 2:
                styles_used[style_name]['examples'].append(para.text[:50])

            # 提取格式信息
            if para.runs:
                run = para.runs[0]
                if run.font.name:
                    styles_used[style_name]['font_name'] = run.font.name
                if run.font.size:
                    styles_used[style_name]['font_size'] = run.font.size.pt
                if run.font.bold is not None:
                    styles_used[style_name]['bold'] = run.font.bold

            # 段落格式
            if para.alignment is not None:
                styles_used[style_name]['alignment'] = str(para.alignment)

            if para.paragraph_format.line_spacing is not None:
                styles_used[style_name]['line_spacing'] = para.paragraph_format.line_spacing

            if para.paragraph_format.space_before is not None:
                styles_used[style_name]['space_before'] = para.paragraph_format.space_before.pt

            if para.paragraph_format.space_after is not None:
                styles_used[style_name]['space_after'] = para.paragraph_format.space_after.pt

            if para.paragraph_format.first_line_indent is not None:
                styles_used[style_name]['first_line_indent'] = para.paragraph_format.first_line_indent.cm

    # 打印样式信息
    for style_name, info in sorted(styles_used.items()):
        print(f"\n样式: {style_name}")
        print(f"  示例文本: {info['examples'][0] if info['examples'] else 'N/A'}")
        if info['font_name']:
            print(f"  字体: {info['font_name']}")
        if info['font_size']:
            print(f"  字号: {info['font_size']} pt")
        if info['bold'] is not None:
            print(f"  加粗: {'是' if info['bold'] else '否'}")
        if info['alignment']:
            print(f"  对齐: {info['alignment']}")
        if info['line_spacing']:
            print(f"  行距: {info['line_spacing']}")
        if info['space_before']:
            print(f"  段前: {info['space_before']} pt")
        if info['space_after']:
            print(f"  段后: {info['space_after']} pt")
        if info['first_line_indent']:
            print(f"  首行缩进: {info['first_line_indent']:.2f} cm")

    # 分析表格
    print(f"\n【表格】")
    print(f"表格数量: {len(doc.tables)}")

    print("\n" + "=" * 60)

    return styles_used

if __name__ == "__main__":
    template_file = "2026届论文格式模版（最终版）(1).docx"
    analyze_docx_format(template_file)
