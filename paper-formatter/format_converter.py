#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
论文格式转换器核心模块
根据2026届论文格式模版（最终版）的要求自动调整Word文档格式
"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT, WD_LINE_SPACING
from docx.oxml.ns import qn
import os

class PaperFormatConverter:
    """论文格式转换器"""

    def __init__(self):
        """初始化格式配置"""
        # 页面设置
        self.page_config = {
            'top_margin': Cm(2.5),
            'bottom_margin': Cm(2.5),
            'left_margin': Cm(3.0),
            'right_margin': Cm(3.0),
            'page_width': Cm(21.0),   # A4纸宽度
            'page_height': Cm(29.7),  # A4纸高度
            'footer_distance': Cm(2.0),  # 页脚距边界
        }

        # 字体配置
        self.fonts = {
            'songti': '宋体',
            'heiti': '黑体',
            'times': 'Times New Roman',
        }

        # 样式配置（基于2026届论文格式模版最终版实际分析）
        self.style_config = {
            'heading1': {
                # 一级标题：如"绪论"、"第一章"
                'font_name': '黑体',
                'font_size': Pt(16),
                'bold': True,
                'alignment': WD_PARAGRAPH_ALIGNMENT.CENTER,
                'space_before': Pt(6),
                'space_after': Pt(6),
            },
            'heading2': {
                # 二级标题：如"1.1 研究背景"
                'font_name': '宋体',
                'font_size': Pt(14),
                'bold': True,
                'alignment': WD_PARAGRAPH_ALIGNMENT.LEFT,
                'first_line_indent': Cm(0.99),
                'space_before': Pt(3),
                'space_after': Pt(3),
            },
            'heading3': {
                # 三级标题
                'font_name': '宋体',
                'font_size': Pt(12),
                'bold': True,
                'alignment': WD_PARAGRAPH_ALIGNMENT.LEFT,
            },
            'normal': {
                # 正文段落
                'font_name': '宋体',
                'font_name_ascii': 'Times New Roman',  # 英文使用Times New Roman
                'font_size': Pt(12),
                'alignment': WD_PARAGRAPH_ALIGNMENT.JUSTIFY,  # 两端对齐
                'first_line_indent': Cm(0.80),  # 首行缩进约2字符
                'line_spacing': 1.5,  # 1.5倍行距
            },
            'caption': {
                # 图表标题：如"表2-1 用户表表结构"
                'font_name': '宋体',
                'font_size': Pt(10.5),
                'alignment': WD_PARAGRAPH_ALIGNMENT.LEFT,
                'first_line_indent': Cm(0.64),  # 图表标题有首行缩进
                'space_after': Pt(6),
            },
            'cover_title': {
                # 封面标题："本科毕业论文"
                'font_name': '黑体',
                'font_size': Pt(42),
                'bold': True,
                'alignment': WD_PARAGRAPH_ALIGNMENT.CENTER,
                'line_spacing': 1.0,
            },
            'paper_title': {
                # 论文题目
                'font_name': '宋体',
                'font_size': Pt(16),
                'bold': True,
                'alignment': WD_PARAGRAPH_ALIGNMENT.CENTER,
                'line_spacing': 1.5,
            }
        }

    def convert(self, input_path, output_path):
        """
        转换文档格式

        Args:
            input_path: 输入文档路径
            output_path: 输出文档路径

        Returns:
            (success, message): 成功标志和消息
        """
        try:
            # 加载文档
            doc = Document(input_path)

            # 应用页面设置
            self._apply_page_setup(doc)

            # 应用段落格式
            self._apply_paragraph_formats(doc)

            # 应用表格格式
            self._apply_table_formats(doc)

            # 保存文档
            doc.save(output_path)

            return True, f"格式转换成功！已保存为: {os.path.basename(output_path)}"

        except Exception as e:
            return False, f"格式转换失败: {str(e)}"

    def _apply_page_setup(self, doc):
        """应用页面设置"""
        for section in doc.sections:
            section.top_margin = self.page_config['top_margin']
            section.bottom_margin = self.page_config['bottom_margin']
            section.left_margin = self.page_config['left_margin']
            section.right_margin = self.page_config['right_margin']
            section.page_width = self.page_config['page_width']
            section.page_height = self.page_config['page_height']
            section.footer_distance = self.page_config['footer_distance']

    def _apply_paragraph_formats(self, doc):
        """应用段落格式"""
        for para in doc.paragraphs:
            if not para.text.strip():
                continue

            text = para.text.strip()
            style_name = para.style.name.lower()

            # 封面标题："本科毕业论文"或"本科毕业设计"
            if text in ('本科毕业论文', '本科毕业设计') and len(text) <= 10:
                self._apply_style(para, 'cover_title')
            # 论文题目（包含"题目："的段落）
            elif text.startswith('题目：') or text.startswith('题目:'):
                self._apply_style(para, 'paper_title')
            # 一级标题：Heading 1样式或章节标题
            elif 'heading 1' in style_name or text in ('绪论', '摘要', 'Abstract', '参考文献', '致谢'):
                self._apply_style(para, 'heading1')
            elif (text.startswith('第') and ('章' in text or '节' in text)) and len(text) <= 15:
                self._apply_style(para, 'heading1')
            # 二级标题：Heading 2样式
            elif 'heading 2' in style_name:
                self._apply_style(para, 'heading2')
            # 三级标题：Heading 3样式
            elif 'heading 3' in style_name:
                self._apply_style(para, 'heading3')
            # 图表标题：以"图"或"表"开头
            elif 'caption' in style_name or (text.startswith(('图', '表')) and len(text) < 50):
                self._apply_style(para, 'caption')
            else:
                # 正文段落
                self._apply_style(para, 'normal')

    def _apply_style(self, para, style_type):
        """
        应用指定样式到段落

        Args:
            para: 段落对象
            style_type: 样式类型 (heading1, heading2, heading3, normal, caption)
        """
        config = self.style_config.get(style_type, {})

        # 段落格式
        if 'alignment' in config:
            para.alignment = config['alignment']

        if 'first_line_indent' in config:
            para.paragraph_format.first_line_indent = config['first_line_indent']

        if 'space_before' in config:
            para.paragraph_format.space_before = config['space_before']

        if 'space_after' in config:
            para.paragraph_format.space_after = config['space_after']

        if 'line_spacing' in config:
            para.paragraph_format.line_spacing = config['line_spacing']

        # 字体格式
        for run in para.runs:
            if 'font_name' in config:
                run.font.name = config['font_name']
                # 设置中文字体
                run._element.rPr.rFonts.set(qn('w:eastAsia'), config['font_name'])

            if 'font_name_ascii' in config:
                # 设置ASCII字符字体
                run._element.rPr.rFonts.set(qn('w:ascii'), config['font_name_ascii'])
                run._element.rPr.rFonts.set(qn('w:hAnsi'), config['font_name_ascii'])

            if 'font_size' in config:
                run.font.size = config['font_size']

            if 'bold' in config:
                run.font.bold = config['bold']

    def _apply_table_formats(self, doc):
        """应用表格格式"""
        for table in doc.tables:
            # 设置表格样式
            table.style = 'Table Grid'

            # 设置表格中的文字格式
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        # 表格内容使用宋体
                        for run in para.runs:
                            run.font.name = '宋体'
                            run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
                            # 数据表格使用10.5pt，信息表格使用14pt
                            # 这里统一使用10.5pt（五号字）
                            run.font.size = Pt(10.5)

                        # 表格内容居中对齐
                        para.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    def get_format_summary(self):
        """获取格式要求摘要（基于2026届论文格式模版最终版）"""
        summary = {
            '页面设置': {
                '纸张': 'A4 (21.0 x 29.7 cm)',
                '上边距': '2.5 cm',
                '下边距': '2.5 cm',
                '左边距': '3.0 cm',
                '右边距': '3.0 cm',
                '页脚距边界': '2.0 cm',
            },
            '封面': {
                '本科毕业论文': '黑体 42pt 居中',
                '论文题目': '宋体 16pt 加粗 居中 1.5倍行距',
            },
            '一级标题': '黑体 16pt 居中 段前段后6pt',
            '二级标题': '宋体 14pt 加粗 左对齐 首行缩进0.99cm',
            '三级标题': '宋体 12pt 加粗 左对齐',
            '正文': '宋体/Times New Roman 12pt 两端对齐 首行缩进0.80cm 1.5倍行距',
            '图表标题': '宋体 10.5pt 左对齐 首行缩进0.64cm',
            '表格': '宋体 10.5pt 居中对齐',
        }
        return summary
